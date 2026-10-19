"""Los datos de entrada del estudio del voto CERA (`data/cera/`)."""
import glob
import json
import os
import shutil

import openpyxl
import pandas as pd
import pytest

from mtpy.core.io import FileSystem
from mtpy.lib.studies.cera.data import (
    MONTHLY_COLUMNS,
    national_series,
    normalise_label,
    read_candidatures_map,
    read_election_census,
    read_monthly,
    read_monthly_workbook,
    read_provinces,
    read_results,
    read_scrutiny,
    scrutiny_votes
)

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))


def test_cera_data_folder_is_complete():
    root = os.path.join(ROOT, 'data', 'cera')
    assert len(glob.glob(os.path.join(root, 'ine', 'cerapropais_*.xls*'))) == 21
    assert all(os.path.exists(os.path.join(root, 'ine', f'elecgen2026_tab{n}.xlsx')) for n in (6, 7, 8))
    params = json.load(open(os.path.join(root, 'params.json'), encoding='utf-8'))
    assert params['election']['date'] == '2026-11-29' and params['check']['max_gain'] == 1
    assert params['participation']['latam'] == 0.05 and params['litigation']['count'] == 163020
    countries = pd.read_csv(os.path.join(root, 'countries.csv'))
    assert set(countries['region']) <= {'latam', 'europe', 'usa_canada', 'other'} and set(countries['exposure']) == {'alta', 'media', 'baja'}
    provinces = json.load(open(os.path.join(root, 'provinces.json'), encoding='utf-8'))
    assert set(provinces.values()) == set(range(1, 53))
    for date in ('2019-11-10', '2023-07-23'):
        frame = pd.read_csv(os.path.join(root, 'escrutinio', f'es-{date}.csv'))
        assert sorted(frame['region_id']) == list(range(1, 53)) and {'electors', 'voters', 'invalid', 'blank', 'PP', 'PSOE'} <= set(frame.columns)
    assert frame.set_index('region_id').loc[28, 'voters'] == 44807  # Madrid 2023


@pytest.fixture(scope='module')
def data_fs():
    return FileSystem(os.path.join(ROOT, 'data'))


@pytest.fixture(scope='module')
def provinces(data_fs):
    return read_provinces(data_fs)


def path(name):
    return FileSystem(os.path.join(ROOT, 'data')).get_path(name)


def test_normalise_label_strips_accents_and_punctuation():
    assert normalise_label('C.F. de Navarra') == 'CFDENAVARRA' and normalise_label('Araba/Álava') == 'ARABAALAVA'


def test_read_monthly_workbook_reads_both_formats(data_fs, provinces):
    old = read_monthly_workbook(path('cera/ine/cerapropais_06.xls'), provinces)
    new = read_monthly_workbook(path('cera/ine/cerapropais_26.xlsx'), provinces)
    assert list(old.columns) == MONTHLY_COLUMNS and old['date'].nunique() == 12 and new['date'].nunique() == 9
    assert national_series(old)['2006-12-01'] == 1160209 and national_series(new)['2026-05-01'] == 2708083
    assert national_series(new)['2026-09-01'] == 2762829 and set(new['region_id']) == set(range(1, 53))


def _albacete_cell(sheet, country, month):
    """(fila, columna) en base 1 del país `country` de Albacete en la columna del mes `month`."""
    column = next(
        cell.column for row in sheet.iter_rows(max_row=10) for cell in row
        if isinstance(cell.value, str) and cell.value.strip().endswith(month)
    )
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    start = next(i for i, values in enumerate(rows) if str(values[0]).strip().upper() == 'TOTAL ALBACETE')
    offset = next(i for i, values in enumerate(rows[start + 1:]) if str(values[2]).strip() == country)
    return start + 2 + offset, column


def test_read_monthly_keeps_the_newest_reading_of_a_month(tmp_path, provinces):
    # dos libros con el mismo mes y cifras distintas: manda el último por nombre de fichero
    folder = tmp_path / 'cera' / 'ine'
    folder.mkdir(parents=True)
    shutil.copy(path('cera/ine/cerapropais_26.xlsx'), folder / 'cerapropais_26.xlsx')
    book = openpyxl.load_workbook(folder / 'cerapropais_26.xlsx')
    sheet = book[book.sheetnames[0]]
    row, column = _albacete_cell(sheet, 'ALEMANIA', 'a 1-ene-2026')
    assert sheet.cell(row=row, column=column).value == 610
    sheet.cell(row=row, column=column).value = 999
    book.save(folder / 'cerapropais_27.xlsx')
    long = read_monthly(FileSystem(str(tmp_path)), provinces)
    selected = (long['date'] == '2026-01-01') & (long['region_id'] == 2) & (long['country'] == 'ALEMANIA')
    assert long[selected]['electors'].item() == 999
    assert len(long) == len(read_monthly_workbook(str(folder / 'cerapropais_26.xlsx'), provinces))


def test_read_monthly_series_has_245_months_to_may_2026(data_fs, provinces):
    series = national_series(read_monthly(data_fs, provinces))
    assert len(series[:'2026-05-01']) == 245 and series.index.min() == pd.Timestamp('2006-01-01')


def test_read_election_census_totals(data_fs, provinces):
    census = read_election_census(path('cera/ine/elecgen2026_tab8.xlsx'), provinces)
    assert census['electors'].sum() == 2769645 and set(census['region_id']) == set(range(1, 53))
    assert census[census['region_id'] == 28]['electors'].sum() == 498874


def test_read_results_2023_matches_the_official_totals(data_fs):
    results = read_results(data_fs, '2023-07-23')
    assert results.seats.values.sum() == 350 and results.totals['registered_cera'].sum() == 2328260
    assert results.votes.loc[28, 'PARTIDO POPULAR'] > 1_000_000 and results.siglas['PARTIDO POPULAR'] == 'PP'


def test_scrutiny_votes_maps_columns_and_rejects_unknown_ones(data_fs):
    scrutiny = read_scrutiny(path('cera/escrutinio/es-2023-07-23.csv'))
    votes = scrutiny_votes(scrutiny, read_candidatures_map(data_fs)['2023-07-23'])
    assert votes.loc[28, 'PARTIDO POPULAR'] == scrutiny.loc[28, 'PP'] and votes.values.sum() == 203152
    assert scrutiny.loc[28, 'voters'] == 44807
    with pytest.raises(ValueError, match='NUEVA'):
        scrutiny_votes(scrutiny.assign(NUEVA=1), read_candidatures_map(data_fs)['2023-07-23'])
